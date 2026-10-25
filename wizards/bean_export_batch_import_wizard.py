# -*- coding: utf-8 -*-
"""Import Excel des lots d'usinage (historique et saisie en masse).

Les colonnes sont reconnues par leur EN-TÊTE, avec des alias (« N° lot
site », « Lot », « Lot usinage »…), dans n'importe quel ordre : un classeur
existant s'importe sans être remis en forme. Un modèle vierge est proposé.

Deux temps : « Analyser » contrôle chaque ligne sans rien créer (site,
certification, poids, doublons dans le fichier et en base) ; « Importer »
crée les seules lignes valides, en brouillon ou confirmées.
"""

import base64
import io
import json
import re
import unicodedata
from datetime import date, datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover - dépendance du socle
    openpyxl = None

# Champ → (libellé du modèle, alias reconnus, obligatoire)
COLUMNS = [
    ('site_lot_ref', "N° lot site", ('n lot site', 'no lot site', 'lot site',
                                     'lot usinage', 'n lot', 'lot', 'lots'), True),
    ('weighing_ticket', "N° bulletin de pesée", ('n bulletin', 'bulletin',
                                                  'bulletin de pesee', 'ticket',
                                                  'n ticket', 'n bulletin de pesee'), False),
    ('site', "Site d'usinage", ('site', 'site usinage', 'usine', 'magasin'), True),
    ('date', "Date", ('date', 'date pesee', 'date de pesee'), True),
    ('bag_count', "Sacs", ('sacs', 'nombre de sacs', 'nb sacs', 'nbre sacs'), False),
    ('gross_weight_kg', "Poids brut (kg)", ('poids brut', 'brut', 'poids brut kg'), True),
    ('tare_weight_kg', "Tare (kg)", ('tare', 'tare kg'), False),
    ('net_weight_kg', "Poids net (kg)", ('poids net', 'net', 'poids net kg'), False),
    ('certification', "Certification", ('certification', 'certif', 'label'), False),
    ('grade', "Grade", ('grade', 'qualite', 'grade qualite'), False),
    ('moisture_pct', "Humidité (%)", ('humidite', 'humidite %', 'hum'), False),
    ('bean_count', "Grainage", ('grainage', 'grainage feves 100 g'), False),
    ('mouldy_pct', "Moisies (%)", ('moisies', 'moisi', 'moisies %'), False),
    ('slaty_pct', "Ardoisées (%)", ('ardoisees', 'ardoise', 'ardoisees %'), False),
    ('insect_damaged_pct', "Mitées (%)", ('mitees', 'insectes', 'mitees %'), False),
    ('quality_certificate_ref', "N° certificat qualité", ('certificat qualite',
                                                          'n certificat qualite'), False),
    ('note', "Observations", ('observations', 'observation', 'notes', 'note'), False),
]
FLOAT_FIELDS = ('gross_weight_kg', 'tare_weight_kg', 'net_weight_kg', 'moisture_pct',
                'mouldy_pct', 'slaty_pct', 'insect_damaged_pct')
INT_FIELDS = ('bag_count', 'bean_count')
GRADE_ALIASES = {
    'gf': 'GF', 'g1': 'GF', 'grade 1': 'GF', 'grade i': 'GF', 'good fermented': 'GF',
    'f': 'F', 'g2': 'F', 'grade 2': 'F', 'grade ii': 'F', 'fair fermented': 'F',
    'ss': 'SS', 'sg': 'SS', 'sous grade': 'SS', 'sub standard': 'SS',
    'limit': 'LIMIT', 'limite': 'LIMIT',
}


def _norm(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode()
    return ' '.join(re.sub(r'[^a-z0-9%]+', ' ', text.lower()).split())


def _to_float(value):
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace('\xa0', '').replace(' ', '').replace(',', '.')
    return float(text)


def _to_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d.%m.%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


class CocoaBeanExportBatchImportWizard(models.TransientModel):
    _name = 'cocoa.bean.export.batch.import.wizard'
    _description = "Import Excel des lots d'usinage"

    state = fields.Selection([('upload', "Fichier"), ('preview', "Aperçu"), ('done', "Terminé")],
                             default='upload', required=True)
    file_data = fields.Binary(string="Fichier Excel (.xlsx)")
    file_name = fields.Char(string="Nom du fichier")
    template_file = fields.Binary(string="Modèle", readonly=True)
    template_file_name = fields.Char(default="modele_lots_usinage.xlsx")
    confirm_batches = fields.Boolean(
        string="Confirmer les lots importés",
        help="Décoché : les lots sont créés en brouillon, à vérifier puis confirmer.")
    create_missing_sites = fields.Boolean(
        string="Créer les sites inconnus",
        help="Crée un site d'usinage (externe) pour tout nom absent d'Odoo.")
    line_ids = fields.One2many('cocoa.bean.export.batch.import.line', 'wizard_id',
                               string="Lignes")
    valid_count = fields.Integer(compute='_compute_counts')
    error_count = fields.Integer(compute='_compute_counts')
    result = fields.Text(readonly=True)

    @api.depends('line_ids.status')
    def _compute_counts(self):
        for wizard in self:
            wizard.valid_count = len(wizard.line_ids.filtered(lambda l: l.status == 'ok'))
            wizard.error_count = len(wizard.line_ids) - wizard.valid_count

    def _reopen(self):
        return {'type': 'ir.actions.act_window', 'res_model': self._name,
                'res_id': self.id, 'view_mode': 'form', 'target': 'new'}

    # ------------------------------------------------------------------
    # Modèle
    # ------------------------------------------------------------------
    def action_download_template(self):
        self.ensure_one()
        if not openpyxl:
            raise UserError(_("La bibliothèque openpyxl est requise."))
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Lots d'usinage"
        sheet.append([label for _f, label, _a, _r in COLUMNS])
        sheet.append(['5011', 'DS2627U00763', 'SITAPA', '24/10/2026', 400, 26400, 400,
                      None, 'RA', 'GF', 7.2, 105, 2.5, 3.0, 0.5, 'QC-0001', ''])
        for index, (_f, label, _a, required) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='5D4037' if required else '8D6E63')
            sheet.column_dimensions[cell.column_letter].width = max(14, len(label) + 2)
        notes = workbook.create_sheet("Instructions")
        notes.append(["Colonne", "Obligatoire", "Contenu"])
        explanations = {
            'net_weight_kg': "Facultatif : contrôlé contre brut − tare s'il est saisi.",
            'certification': "Suffixe, libellé commercial ou nom de la certification ; vide = conventionnel.",
            'grade': "GF / F / SS / Limite (ou Grade 1, Grade 2, sous-grade).",
            'date': "jj/mm/aaaa ou date Excel.",
        }
        for field_name, label, _a, required in COLUMNS:
            notes.append([label, "oui" if required else "", explanations.get(field_name, '')])
        buffer = io.BytesIO()
        workbook.save(buffer)
        self.template_file = base64.b64encode(buffer.getvalue())
        return {'type': 'ir.actions.act_url', 'target': 'self',
                'url': '/web/content/%s/%s/template_file/%s?download=true' % (
                    self._name, self.id, self.template_file_name)}

    # ------------------------------------------------------------------
    # Analyse
    # ------------------------------------------------------------------
    def _read_rows(self):
        if not openpyxl:
            raise UserError(_("La bibliothèque openpyxl est requise."))
        if not self.file_data:
            raise UserError(_("Choisissez un fichier Excel."))
        if self.file_name and not self.file_name.lower().endswith('.xlsx'):
            raise UserError(_("Le fichier doit être au format .xlsx (reçu : %s).", self.file_name))
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.file_data)), data_only=True, read_only=True)
        except Exception as exc:  # fichier corrompu ou non Excel
            raise UserError(_("Fichier illisible : %s", exc)) from None
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError(_("La première feuille est vide."))
        header = [_norm(cell) for cell in rows[0]]
        mapping = {}
        for field_name, label, aliases, _required in COLUMNS:
            for candidate in (_norm(label),) + aliases:
                if candidate in header:
                    mapping[field_name] = header.index(candidate)
                    break
        missing = [label for field_name, label, _a, required in COLUMNS
                   if required and field_name not in mapping]
        if missing:
            raise UserError(_(
                "Colonnes obligatoires introuvables : %(missing)s.\n\nEn-têtes lus : "
                "%(found)s", missing=", ".join(missing),
                found=", ".join(str(c) for c in rows[0] if c)))
        for number, row in enumerate(rows[1:], start=2):
            if not any(cell not in (None, '') for cell in row):
                continue
            yield number, {field_name: row[index] if index < len(row) else None
                           for field_name, index in mapping.items()}

    def _find_site(self, name):
        Site = self.env['cocoa.reception.site']
        return Site.search([('name', '=ilike', name.strip()),
                            ('company_id', '=', self.env.company.id)], limit=1)

    def _find_certification(self, text):
        Cert = self.env['potting.certification']
        text = text.strip()
        # Suffixe (code), libellé commercial (« RA-SAS ») ou nom complet
        return (Cert.search([('code', '=ilike', text)], limit=1)
                or Cert.search([('premium_label', '=ilike', text)], limit=1)
                or Cert.search([('name', '=ilike', text)], limit=1))

    def action_parse_file(self):
        self.ensure_one()
        self.line_ids.unlink()
        Batch = self.env['cocoa.bean.export.batch']
        seen = set()
        lines = []
        for number, raw in self._read_rows():
            errors, vals = [], {}
            text = {k: ('' if v is None else str(v).strip()) for k, v in raw.items()}
            vals['site_lot_ref'] = text.get('site_lot_ref') or False
            if not vals['site_lot_ref']:
                errors.append(_("n° de lot du site manquant"))
            for key in ('weighing_ticket', 'quality_certificate_ref', 'note'):
                vals[key] = text.get(key) or False
            site_name = text.get('site', '')
            site = self._find_site(site_name) if site_name else self.env['cocoa.reception.site']
            if not site_name:
                errors.append(_("site manquant"))
            elif not site and not self.create_missing_sites:
                errors.append(_("site « %s » inconnu", site_name))
            try:
                vals['date'] = _to_date(raw.get('date'))
                if not vals['date']:
                    errors.append(_("date manquante"))
            except ValueError:
                errors.append(_("date illisible « %s »", text.get('date')))
            for key in FLOAT_FIELDS + INT_FIELDS:
                if key not in raw:
                    continue
                try:
                    value = _to_float(raw[key])
                except ValueError:
                    errors.append(_("nombre illisible « %(v)s » (%(c)s)", v=text.get(key), c=key))
                    continue
                if value is not None:
                    vals[key] = int(round(value)) if key in INT_FIELDS else value
            gross, tare = vals.get('gross_weight_kg'), vals.get('tare_weight_kg') or 0.0
            if not gross or gross <= 0:
                errors.append(_("poids brut manquant ou nul"))
            elif tare >= gross:
                errors.append(_("tare ≥ poids brut"))
            net = vals.pop('net_weight_kg', None)
            if net is not None and gross and abs((gross - tare) - net) > 0.5:
                errors.append(_("net %(net)s ≠ brut − tare (%(calc)s)", net=net,
                                calc=round(gross - tare, 2)))
            cert_text = text.get('certification')
            if cert_text:
                cert = self._find_certification(cert_text)
                if cert:
                    vals['certification_id'] = cert.id
                else:
                    errors.append(_("certification « %s » inconnue", cert_text))
            if text.get('grade'):
                grade = GRADE_ALIASES.get(_norm(text['grade']))
                if grade:
                    vals['grade'] = grade
                else:
                    errors.append(_("grade « %s » inconnu", text['grade']))
            # Doublons : même site + même bulletin (ou lot si pas de bulletin)
            key = (_norm(site_name), _norm(vals.get('weighing_ticket') or vals['site_lot_ref']))
            if key in seen:
                errors.append(_("doublon dans le fichier"))
            seen.add(key)
            if site and vals.get('weighing_ticket') and Batch.search_count([
                    ('site_id', '=', site.id), ('state', '!=', 'cancelled'),
                    ('weighing_ticket', '=ilike', vals['weighing_ticket'])]):
                errors.append(_("bulletin déjà saisi dans Odoo"))
            lines.append({
                'row_number': number,
                'site_name': site_name,
                'site_id': site.id if site else False,
                'site_lot_ref': vals['site_lot_ref'],
                'weighing_ticket': vals.get('weighing_ticket'),
                'net_weight_kg': (gross or 0.0) - tare if gross else 0.0,
                'bag_count': vals.get('bag_count') or 0,
                'values': json.dumps({k: (v.isoformat() if isinstance(v, date) else v)
                                      for k, v in vals.items()}),
                'status': 'error' if errors else 'ok',
                'errors': "; ".join(errors) or False,
            })
        if not lines:
            raise UserError(_("Aucune ligne de données dans le fichier."))
        self.write({'line_ids': [(0, 0, line) for line in lines], 'state': 'preview'})
        return self._reopen()

    # ------------------------------------------------------------------
    # Import
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        valid = self.line_ids.filtered(lambda l: l.status == 'ok')
        if not valid:
            raise UserError(_("Aucune ligne valide à importer."))
        Batch = self.env['cocoa.bean.export.batch']
        Site = self.env['cocoa.reception.site']
        created = Batch
        for line in valid:
            vals = json.loads(line.values)
            if isinstance(vals.get('date'), str):
                vals['date'] = fields.Date.to_date(vals['date'])
            site = line.site_id or self._find_site(line.site_name)
            if not site:
                site = Site.create({'name': line.site_name.strip(), 'site_type': 'external'})
            vals['site_id'] = site.id
            created |= Batch.create(vals)
        if self.confirm_batches:
            created.action_confirm()
        self.write({
            'state': 'done',
            'result': _("%(count)s lot(s) d'usinage importé(s)%(state)s ; %(errors)s ligne(s) "
                        "en erreur ignorée(s).", count=len(created),
                        state=_(" et confirmé(s)") if self.confirm_batches else _(" en brouillon"),
                        errors=self.error_count),
        })
        return {
            'type': 'ir.actions.act_window',
            'name': _("Lots d'usinage importés"),
            'res_model': 'cocoa.bean.export.batch',
            'view_mode': 'tree,form',
            'domain': [('id', 'in', created.ids)],
        }


class CocoaBeanExportBatchImportLine(models.TransientModel):
    _name = 'cocoa.bean.export.batch.import.line'
    _description = "Ligne d'import de lot d'usinage"
    _order = 'row_number'

    wizard_id = fields.Many2one('cocoa.bean.export.batch.import.wizard', required=True,
                                ondelete='cascade')
    row_number = fields.Integer(string="Ligne")
    site_name = fields.Char(string="Site (fichier)")
    site_id = fields.Many2one('cocoa.reception.site', string="Site")
    site_lot_ref = fields.Char(string="Lot site")
    weighing_ticket = fields.Char(string="Bulletin")
    bag_count = fields.Integer(string="Sacs")
    net_weight_kg = fields.Float(string="Net (kg)", digits=(16, 2))
    values = fields.Text(string="Valeurs")
    status = fields.Selection([('ok', "Valide"), ('error', "Erreur")], string="Statut")
    errors = fields.Text(string="Erreurs")
